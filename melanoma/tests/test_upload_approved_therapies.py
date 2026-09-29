"""Sheet rules for the approved-therapies loader.

Built from an in-memory workbook shaped like the real one: row 1's cancer-type
headers merged over each US/EU pair, row 2 reading US/EU, then section header
rows and regimen rows. Every EU cell reads "Approved", so a test sees at once
if an EU value leaks into a row.
"""

from __future__ import annotations

from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from scripts.upload_approved_therapies import (
    CANCER_TYPES,
    LoadError,
    build_rows,
    count_by_cancer_type,
)

HEADERS = list(CANCER_TYPES)
DASH = "\u2014"
SOURCE = "NCCN v3.2026"


def _regimen(
    name: str,
    setting: str = "1L+ Advanced",
    us: dict[str, str] | None = None,
    *,
    default: str = DASH,
    dose: str = "200 mg q3w",
) -> list[Any]:
    """One regimen row. `us` maps a sheet header to that column's US status."""
    us = us or {}
    cells: list[Any] = []
    for header in HEADERS:
        cells += [us.get(header, default), "Approved"]
    return [name, None, setting, "Preferred", "Cat 1", dose, *cells]


def _sheet(*rows: list[Any], headers: list[str] = HEADERS) -> Worksheet:
    ws = Workbook().active
    ws.title = "Therapy Matrix"
    ws.append(
        [
            "Agent / Regimen",
            "Biomarker",
            "Treatment setting",
            "NCCN v3.2026 tier",
            "NCCN category",
            "Approved dose / schedule",
            *[value for header in headers for value in (header, None)],
        ]
    )
    for i in range(len(headers)):
        ws.merge_cells(
            start_row=1, start_column=7 + 2 * i, end_row=1, end_column=8 + 2 * i
        )
    ws.append([None] * 6 + ["US", "EU"] * len(headers))
    for row in rows:
        ws.append(row)
    return ws


# Every cancer type needs at least one row, or the loader refuses.
EVERYWHERE = _regimen("Nivolumab", default="On-label (generic)")


def test_merged_header_cells_right_of_the_us_column_read_empty() -> None:
    # The assumption the loader is built on: openpyxl gives a merged range's
    # value to its top-left cell only.
    ws = _sheet(EVERYWHERE)
    assert ws.cell(row=1, column=7).value == HEADERS[0]
    assert ws.cell(row=1, column=8).value is None


def test_reads_each_us_column_under_its_merged_header() -> None:
    rows = build_rows(_sheet(EVERYWHERE), SOURCE)
    assert sorted(r["cancer_type"] for r in rows) == sorted(CANCER_TYPES.values())
    assert {r["us_status"] for r in rows} == {"On-label (generic)"}


def test_dash_means_no_row_in_status_and_null_elsewhere() -> None:
    ws = _sheet(
        EVERYWHERE, _regimen("Tebentafusp", "1L+", {"Uveal Mel": "Approved"}, dose=DASH)
    )
    teben = [r for r in build_rows(ws, SOURCE) if r["treatment_name"] == "Tebentafusp"]
    assert [r["cancer_type"] for r in teben] == ["Uveal Melanoma"]
    assert teben[0]["dose"] is None


def test_skips_section_headers_and_blank_rows() -> None:
    blank = [None] * (6 + 2 * len(HEADERS))
    ws = _sheet(["ADJUVANT SYSTEMIC THERAPY (NCCN ADJSYS-1)"], EVERYWHERE, blank, blank)
    rows = build_rows(ws, SOURCE)
    assert len(rows) == len(CANCER_TYPES)
    assert {r["sheet_row"] for r in rows} == {4}


def test_strips_whitespace_and_stamps_source_and_sheet_row() -> None:
    rows = build_rows(
        _sheet(_regimen("Nivolumab ", default=" On-label (generic) ")), SOURCE
    )
    assert {r["treatment_name"] for r in rows} == {"Nivolumab"}
    assert {r["us_status"] for r in rows} == {"On-label (generic)"}
    assert {(r["source"], r["sheet_row"]) for r in rows} == {(SOURCE, 3)}


def test_keeps_a_status_wording_it_has_never_seen() -> None:
    status = "On-label; NCCN-preferred for asymptomatic brain mets"
    ws = _sheet(
        EVERYWHERE, _regimen("Ipilimumab + Nivolumab", us={"CNS / Brain Mets": status})
    )
    assert status in {r["us_status"] for r in build_rows(ws, SOURCE)}


def test_counts_rows_per_cancer_type() -> None:
    ws = _sheet(EVERYWHERE, _regimen("Tebentafusp", "1L+", {"Uveal Mel": "Approved"}))
    counts = count_by_cancer_type(build_rows(ws, SOURCE))
    assert counts["Uveal Melanoma"] == 2
    assert counts["Merkel Cell Carcinoma"] == 1


def test_refuses_an_unknown_cancer_type_header() -> None:
    with pytest.raises(LoadError, match="unknown cancer-type header 'Ocular'"):
        build_rows(_sheet(EVERYWHERE, headers=[*HEADERS[:-1], "Ocular"]), SOURCE)


def test_refuses_a_cancer_type_header_that_appears_twice() -> None:
    with pytest.raises(LoadError, match="appears twice"):
        build_rows(_sheet(EVERYWHERE, headers=[*HEADERS[:-1], HEADERS[0]]), SOURCE)


def test_refuses_a_cancer_type_with_no_rows() -> None:
    only_cutaneous = _regimen("Nivolumab", us={"Cutaneous / Met Mel": "Approved"})
    with pytest.raises(LoadError, match="no rows for: .*Merkel Cell Carcinoma"):
        build_rows(_sheet(only_cutaneous), SOURCE)


def test_refuses_a_regimen_row_with_no_name() -> None:
    with pytest.raises(LoadError, match="row 4: has a setting but no regimen name"):
        build_rows(_sheet(EVERYWHERE, _regimen("", default="Approved")), SOURCE)


def test_refuses_a_row_with_statuses_but_no_setting() -> None:
    # A blank setting marks a section header; a row that still carries statuses
    # is a regimen whose setting was cleared, and would silently vanish.
    no_setting = _regimen("Nivolumab + relatlimab", default="On-label (generic)")
    no_setting[2] = None
    with pytest.raises(LoadError, match="row 4: has US statuses but no setting"):
        build_rows(_sheet(EVERYWHERE, no_setting), SOURCE)
