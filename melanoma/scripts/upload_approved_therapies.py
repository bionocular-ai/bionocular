#!/usr/bin/env python3
"""Load the NCCN approved-therapies spreadsheet into Supabase approved_therapies.

Reads the US columns of the "Therapy Matrix" sheet and replaces every row of
the table in one transaction (the replace_approved_therapies database
function), so a failed load leaves the previous version in place. Re-run it
for each NCCN version.

Usage:
    poetry run python3 scripts/upload_approved_therapies.py <xlsx> --source "NCCN v3.2026" [--dry-run]
"""

from __future__ import annotations

import argparse
import os
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from supabase import Client, create_client

_MELANOMA_ROOT = Path(__file__).resolve().parent.parent
SHEET = "Therapy Matrix"
TABLE = "approved_therapies"
FIRST_DATA_ROW = 3
NOT_APPLICABLE = "\u2014"  # the sheet's em dash

# Spreadsheet header -> the cancer_type the web app queries with
# (getDbCancerType in web/src/lib/api.ts). Sheet order.
CANCER_TYPES: dict[str, str] = {
    "Cutaneous / Met Mel": "Cutaneous Melanoma",
    "Acral Mel": "Acral Melanoma",
    "Mucosal Mel": "Mucosal Melanoma",
    "CNS / Brain Mets": "Cutaneous Melanoma with Brain/CNS Metastasis",
    "Uveal Mel": "Uveal Melanoma",
    "BCC": "Basal Cell Carcinoma",
    "cSCC": "Cutaneous Squamous Cell Carcinoma",
    "MCC": "Merkel Cell Carcinoma",
}

# Columns A-F: column name -> 1-based sheet column.
_FIELDS = {
    "treatment_name": 1,
    "biomarker": 2,
    "setting": 3,
    "nccn_tier": 4,
    "nccn_category": 5,
    "dose": 6,
}


class LoadError(Exception):
    """The sheet does not have the shape this loader writes from."""


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return None if value in ("", NOT_APPLICABLE) else value
    return value


def us_columns(ws: Worksheet) -> dict[int, str]:
    """Column index -> cancer_type, for every column whose row 2 reads US.

    Row 1's headers are merged over each US/EU pair, and openpyxl gives a
    merged range's value only to its top-left cell - the US column. Reading
    the header from the US column itself means no column position is assumed.
    """
    found: dict[int, str] = {}
    for col in range(1, ws.max_column + 1):
        if _clean(ws.cell(row=2, column=col).value) != "US":
            continue
        header = _clean(ws.cell(row=1, column=col).value)
        letter = get_column_letter(col)
        if header not in CANCER_TYPES:
            raise LoadError(f"column {letter}: unknown cancer-type header {header!r}")
        if CANCER_TYPES[header] in found.values():
            raise LoadError(f"column {letter}: {header!r} appears twice")
        found[col] = CANCER_TYPES[header]
    missing = sorted(set(CANCER_TYPES.values()) - set(found.values()))
    if missing:
        raise LoadError(f"no US column for: {', '.join(missing)}")
    return found


def count_by_cancer_type(rows: list[dict[str, Any]]) -> Counter[str]:
    return Counter(row["cancer_type"] for row in rows)


def build_rows(ws: Worksheet, source: str) -> list[dict[str, Any]]:
    """One row per regimen per cancer type whose US cell is not the dash."""
    columns = us_columns(ws)
    rows: list[dict[str, Any]] = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        fields = {
            name: _clean(ws.cell(row=r, column=c).value) for name, c in _FIELDS.items()
        }
        if fields["setting"] is None:
            # A section header ("ADJUVANT SYSTEMIC THERAPY ...") or a blank row.
            # A row that still has statuses is a regimen whose setting was
            # cleared, and skipping it would drop it silently.
            if any(_clean(ws.cell(row=r, column=col).value) for col in columns):
                raise LoadError(f"row {r}: has US statuses but no setting")
            continue
        if fields["treatment_name"] is None:
            raise LoadError(f"row {r}: has a setting but no regimen name")
        for col, cancer_type in columns.items():
            status = _clean(ws.cell(row=r, column=col).value)
            if status is None:
                continue
            rows.append(
                {
                    "cancer_type": cancer_type,
                    "sheet_row": r,
                    **fields,
                    "us_status": status,
                    "source": source,
                }
            )
    counts = count_by_cancer_type(rows)
    empty = sorted(t for t in CANCER_TYPES.values() if counts[t] == 0)
    if empty:
        raise LoadError(f"no rows for: {', '.join(empty)}")
    return rows


def _client() -> Client:
    load_dotenv(_MELANOMA_ROOT / ".env")
    url = os.environ.get("SUPABASE_URL")
    # replace_approved_therapies is granted to service_role only.
    key = os.environ.get("SUPABASE_SECRET_KEY")
    if not url or not key:
        sys.exit("SUPABASE_URL and SUPABASE_SECRET_KEY must be set in melanoma/.env")
    return create_client(url, key)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--source", required=True, help='NCCN version, e.g. "NCCN v3.2026"'
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="parse and print counts; write nothing"
    )
    args = parser.parse_args()

    ws = load_workbook(args.workbook)[SHEET]
    try:
        rows = build_rows(ws, args.source)
    except LoadError as error:
        sys.exit(f"refusing to load: {error}")

    for cancer_type, n in sorted(count_by_cancer_type(rows).items()):
        print(f"{n:4d}  {cancer_type}")
    print(f"{len(rows):4d}  total ({args.source})")
    if args.dry_run:
        return

    sb = _client()
    sb.rpc("replace_approved_therapies", {"payload": rows}).execute()
    stored = sb.table(TABLE).select("sheet_row", count="exact").execute().count
    if stored != len(rows):
        sys.exit(f"wrote {len(rows)} rows but {TABLE} holds {stored}")
    print(f"replaced {TABLE}: {stored} rows")


if __name__ == "__main__":
    main()
